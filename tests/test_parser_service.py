import pytest
import fitz
from docx import Document
from services.parser_service import ParserService


def test_instancia_parser():
    parser = ParserService()
    assert parser is not None


def test_parse_md(tmp_path):

    caminho_arquivo = tmp_path / "teste.md"

    caminho_arquivo.write_text("# Documento Teste\n\nConteúdo de exemplo.", encoding="utf-8")
    parser = ParserService()
    resultado = parser.parse_md(
        {
            "nome_arquivo": "teste.md",
            "tipo_arquivo": ".md",
            "caminho": str(caminho_arquivo)
        }
    )
    assert resultado is not None
    assert resultado["nome_arquivo"] == "teste.md"
    assert resultado["tipo_arquivo"] == ".md"
    assert resultado["conteudo"] != ""
    assert resultado["quantidade_linhas"] > 0
    assert resultado["quantidade_caracteres"] > 0


def test_parse_docx(tmp_path):
    caminho_arquivo = tmp_path / "teste.docx"

    documento = Document()
    documento.add_paragraph("Primeiro parágrafo de teste.")
    documento.add_paragraph("Segundo parágrafo de teste.")

    documento.save(caminho_arquivo)

    parser = ParserService()

    resultado = parser.parse_docx(
        {
            "nome_arquivo": "teste.docx",
            "tipo_arquivo": ".docx",
            "caminho": str(caminho_arquivo)
        }
    )
    assert resultado is not None
    assert resultado["nome_arquivo"] == "teste.docx"
    assert resultado["tipo_arquivo"] == ".docx"
    assert resultado["conteudo"] != ""
    assert resultado["qtd_paragrafos"] == 2


def test_parse_pdf(tmp_path):

    caminho_arquivo = tmp_path / "teste.pdf"

    pdf = fitz.open()

    pagina = pdf.new_page()

    pagina.insert_text((72, 72), "Documento PDF de Teste")
    pdf.save(str(caminho_arquivo))
    pdf.close()

    parser = ParserService()
    resultado = parser.parse_pdf(
        {
            "nome_arquivo": "teste.pdf",
            "tipo_arquivo": ".pdf",
            "caminho": str(caminho_arquivo)
        }
    )

    assert resultado is not None
    assert resultado["nome_arquivo"] == "teste.pdf"
    assert resultado["tipo_arquivo"] == ".pdf"
    assert resultado["conteudo"] != ""
    assert resultado["numero_paginas"] == 1


def test_parse_document_md(tmp_path):

    caminho_arquivo = tmp_path / "teste.md"
    caminho_arquivo.write_text("# Documento Teste", encoding="utf-8")

    parser = ParserService()

    resultado = parser.parse_document(
        {
            "nome_arquivo": "teste.md",
            "tipo_arquivo": ".md",
            "caminho": str(caminho_arquivo)
        }
    )

    assert resultado is not None
    assert resultado["tipo_arquivo"] == ".md"


def test_tipo_arquivo_invalido():
    parser = ParserService()
    with pytest.raises(ValueError):
        parser.parse_document(
            {
                "nome_arquivo": "teste.txt",
                "tipo_arquivo": ".txt",
                "caminho": "teste.txt"
            }
        )